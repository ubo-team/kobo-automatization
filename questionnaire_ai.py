"""AI helpers for the Gjenero XLS page.

- Converts an unformatted questionnaire (docx, xlsx, pdf) into the tagged line
  format that the page's parser (generate_xlsform) already understands.
- Adds skip logic (the XLSForm `relevant` column) from the questionnaire's
  routing instructions.
- Tests every `relevant` expression offline: syntax, references, choice values,
  and a simulation over answer combinations (never shown / never hides).
- Reviews the filters of an existing XLSForm.
"""

import base64
import itertools
import json
import math
import os
import re
from io import BytesIO

import anthropic
import openpyxl
import pandas as pd
from docx import Document
from docx2python import docx2python

MODEL = "claude-opus-5"          # the page can override it with CLAUDE_MODEL in secrets (e.g. claude-sonnet-5)
FALLBACK_BETA = "server-side-fallback-2026-07-01"
FALLBACK_MODELS = ("claude-opus-5",)
# USD per 1M tokens (input, output), used only for the cost estimate shown in the UI
PRICING = {"claude-opus-5": (5.00, 25.00), "claude-sonnet-5": (2.00, 10.00)}
# Thinking depth per step. Formatting is mostly careful transcription and most of its cost is output
# (thinking is billed as output), so it runs at "medium"; filters need reasoning over the routing, so
# they keep the default "high". Raise CONVERT_EFFORT to "high" if question types or codes come out wrong.
CONVERT_EFFORT = "medium"
FILTER_EFFORT = "high"

SOURCE_TYPES = ["docx", "xlsx", "pdf", "txt", "csv", "md"]


class AIError(Exception):
    """Error with a message that can be shown to the user as-is."""


# ---------------------------------------------------------------------------
# Claude calls
# ---------------------------------------------------------------------------

def make_client(api_key):
    return anthropic.Anthropic(api_key=api_key)


def estimate_cost(usage):
    price_in, price_out = PRICING.get(MODEL, PRICING["claude-opus-5"])
    # cache writes (5-minute TTL) cost 1.25x the input price, cache reads 0.1x
    input_cost = (usage.input_tokens or 0) * price_in \
        + (usage.cache_creation_input_tokens or 0) * price_in * 1.25 \
        + (usage.cache_read_input_tokens or 0) * price_in * 0.1
    return (input_cost + (usage.output_tokens or 0) * price_out) / 1_000_000


def _call_claude(client, system, content, schema, max_tokens=64000, effort=FILTER_EFFORT):
    """Streams one request with a JSON-schema output and returns (data, usage)."""
    params = dict(
        model=MODEL,
        max_tokens=max_tokens,
        system=system,
        messages=[{"role": "user", "content": content}],
        output_config={"effort": effort, "format": {"type": "json_schema", "schema": schema}},
    )
    if MODEL in FALLBACK_MODELS:
        params.update(betas=[FALLBACK_BETA], fallbacks="default")
    try:
        with client.beta.messages.stream(**params) as stream:
            message = stream.get_final_message()
    except anthropic.AuthenticationError:
        raise AIError("Çelësi API i Claude (ANTHROPIC_API_KEY) nuk është i vlefshëm.")
    except anthropic.RateLimitError:
        raise AIError("Kufiri i kërkesave te Claude u tejkalua. Provo përsëri pas pak minutash.")
    except anthropic.BadRequestError as e:
        if "credit balance" in str(e.message).lower():
            raise AIError("Llogaria e Anthropic nuk ka kredit të mjaftueshëm. Shto kredit te "
                          "console.anthropic.com → Plans & Billing dhe provo përsëri.")
        raise AIError(f"Kërkesa te Claude u refuzua: {e.message}")
    except anthropic.APIStatusError as e:
        raise AIError(f"Gabim nga shërbimi i Claude ({e.status_code}). Provo përsëri.")
    except anthropic.APIConnectionError:
        raise AIError("Nuk u arrit lidhja me Claude. Kontrollo internetin dhe provo përsëri.")

    if message.stop_reason == "refusal":
        raise AIError("Claude refuzoi ta përpunojë këtë dokument.")
    if message.stop_reason == "max_tokens":
        raise AIError("Përgjigja e Claude u ndërpre sepse arriti kufirin e gjatësisë. Pyetësori është shumë i gjatë "
                      "për një përpunim të vetëm: ndaje në pjesë dhe provo përsëri.")

    # If a fallback model took over, only the text after the last switch point is the answer
    blocks = list(message.content)
    start = max((i + 1 for i, b in enumerate(blocks) if b.type == "fallback"), default=0)
    text = "".join(b.text for b in blocks[start:] if b.type == "text")
    try:
        return json.loads(text), message.usage
    except json.JSONDecodeError:
        raise AIError("Përgjigja e Claude nuk ishte në formatin e pritur. Provo përsëri.")


# ---------------------------------------------------------------------------
# Source documents
# ---------------------------------------------------------------------------

# The questionnaire is sent to Claude as compact Markdown: far fewer tokens than the raw document
# (Word tables repeat merged cells, PDFs are billed per page image) with the same content.

def _clean(text):
    text = re.sub(r'[.…]{4,}', ' ', text)            # dotted leaders: "Refused ………… -99"
    text = re.sub(r'[ \t ]+', ' ', text)
    return text.strip()


def _dedupe(lines):
    out = []
    for line in lines:
        if line and (not out or out[-1] != line):
            out.append(line)
    return out


def _docx_table_lines(table):
    """One Markdown line per table row; merged cells (same XML element) appear once."""
    lines = []
    for row in table.rows:
        cells, seen = [], set()
        for cell in row.cells:
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            parts = [_clean(p.text) for p in cell.paragraphs]
            for nested in cell.tables:
                parts.extend(_docx_table_lines(nested))
            text = " <br> ".join(p for p in parts if p)
            if text and text not in cells:
                cells.append(text)
        if len(cells) == 1:
            lines.append(cells[0])
        elif cells:
            lines.append("| " + " | ".join(cells) + " |")
    return lines


def _docx_markdown(data):
    doc = Document(BytesIO(data))
    lines = []
    for block in doc.iter_inner_content():
        if hasattr(block, "rows"):
            lines.extend(_docx_table_lines(block))
            continue
        text = _clean(block.text)
        if not text:
            continue
        style = (block.style.name if block.style is not None else "").lower()
        heading = re.match(r'heading (\d)', style)
        if heading:
            text = "#" * min(int(heading.group(1)), 6) + " " + text
        elif style == "title":
            text = "# " + text
        elif "list" in style:
            text = "- " + text
        lines.append(text)
    markdown = "\n".join(_dedupe(lines))
    # Content outside the body flow (text boxes, some content controls) is not in python-docx's
    # paragraphs; if much of the document's text is missing, use docx2python's full extraction instead
    full = docx2python(BytesIO(data)).text
    full_words, md_words = set(full.split()), set(markdown.split())
    if full_words and len(full_words & md_words) < 0.9 * len(full_words):
        return "\n".join(_dedupe([_clean(l) for l in full.split("\n")]))
    return markdown


def _xlsx_markdown(data):
    wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"## Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [_clean(str(c)) for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(_dedupe(parts))


def _pdf_markdown(data):
    """Text of a PDF, or None when it has no usable text layer (scanned)."""
    try:
        from pypdf import PdfReader
        pages = PdfReader(BytesIO(data)).pages
        texts = [p.extract_text() or "" for p in pages]
    except Exception:
        return None
    if not pages or sum(len(t.strip()) for t in texts) < 200 * len(pages):
        return None
    lines = []
    for n, text in enumerate(texts, 1):
        lines.append(f"<!-- page {n} -->")
        lines.extend(_clean(l) for l in text.split("\n"))
    return "\n".join(_dedupe(lines))


def to_markdown(filename, data):
    """Questionnaire as Markdown text; None for a PDF that must be sent as a document (scanned)."""
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".docx":
        return _docx_markdown(data)
    if ext == ".xlsx":
        return _xlsx_markdown(data)
    if ext == ".pdf":
        return _pdf_markdown(data)
    if ext in (".txt", ".md", ".csv"):
        return data.decode("utf-8", errors="replace")
    raise AIError(f"Formati '{ext}' nuk mbështetet. Përdor .docx, .xlsx, .pdf, .txt ose .csv.")


def load_source(filename, data):
    """Returns Claude content blocks holding the questionnaire (cached across the calls for one file)."""
    text = to_markdown(filename, data)
    if text is None:
        block = {"type": "document",
                 "source": {"type": "base64", "media_type": "application/pdf",
                            "data": base64.standard_b64encode(data).decode("utf-8")}}
    elif not text.strip():
        raise AIError("Dokumenti nuk përmban tekst që mund të lexohet.")
    else:
        block = {"type": "text", "text": f"<questionnaire>\n{text}\n</questionnaire>"}
    return [block]


def _cached(source_blocks):
    """The questionnaire marked for the 5-minute prompt cache: the filter call writes it, the repair call
    that follows right after (same prompt and effort) reads it at a tenth of the price."""
    return [{**b, "cache_control": {"type": "ephemeral"}} for b in source_blocks]


# ---------------------------------------------------------------------------
# Step 1: questionnaire -> tagged lines
# ---------------------------------------------------------------------------

INPUT_NOTE = ("The questionnaire arrives as Markdown converted from the original file: each table row is one "
              "line in the form `| cell | cell |`, `<br>` separates lines inside a cell, and `<!-- page N -->` marks "
              "PDF pages.")

CONVERT_SYSTEM = INPUT_NOTE + """

You convert survey questionnaires into a line-based tagged text format. A strict parser turns these lines into a KoboToolbox XLSForm, so follow the format exactly: it cannot recover from deviations, and a wrong line silently produces a wrong form. The form is used for real fieldwork and analysed by the answer codes, so carry over everything the questionnaire specifies (codes, variable names, languages, types) rather than simplifying it.

## Line format

Output one item per line and never split a question's text across lines.

Question line: `<number>. <question text> <tags>`. Every question line needs exactly one type tag; tags go at the end of the line.

Numbers: copy the questionnaire's question ID exactly as it is written, without changing, normalising or renumbering it: 5.1 stays 5.1, 3a stays 3a, ECE_M01 stays ECE_M01, Q12 stays Q12. The ID is written once, at the start of the line, followed by a dot and a space (`5.1. <text>`), with no space inside the ID. If the questionnaire has no numbering, number the questions 1, 2, 3, … Keep an ID even when another module uses the same one; the program makes the variable names unique. Never leave out a question because its ID or structure is unclear: give it a number (e.g. the previous question's ID plus a letter: ECE1a) and mention it in `notes`.

Type tags:
- `[single]` one answer from a list. The answer options follow, one per line.
- `[multiple]` several answers allowed ("select all that apply"). Options follow, one per line.
- `[text]` open-ended text answer.
- `[numeric]` a whole number (age, count, days, hours, minutes, number of months).
- `[decimal]` a number that can have decimals or be negative: money amounts, profit, income.
- `[date]` a date, or a month and year.
- `[time]` a duration or time of day given in hours and minutes.
- `[scale S(min label)-E(max label)]` a single rating from S to E, e.g. `[scale 1(Aspak dakord)-5(Plotësisht dakord)]`; labels are optional: `[scale 0-10]`. No option lines follow.
- `[matrix single N]` a grid of statements that share the same N answer columns, one answer per row. Next come exactly N lines with the column labels, then one line per row statement. Use `[matrix multiple N]` when a row allows several columns. A table of statements × agreement scale is a matrix. When a series of separately numbered questions shares one answer list (ASSET1A, ASSET1B, …), keep them as separate `[single]` questions so each keeps its own number.
- `[ranking N]` the respondent ranks their top N choices. Options follow, one per line.
- `[note]` text shown without an answer (introduction, consent text, a read-aloud passage, a section introduction). Untagged lines after it are appended to the note until the next tagged line. A note stands exactly where the questionnaire prints it: an introduction goes before the questions it introduces (right after its `[group]` / `[section]` line when it opens a module or section), never after them.
- `[other]` something that should not be coded as a question: interviewer name/ID, date, GPS, start time, respondent name, phone number or address (the form adds these automatically), or items filled in from office records. The line is skipped and listed for the user.

Optional extra tags on question lines:
- `[name: variable_name]` when the questionnaire gives a variable name for the question (e.g. `hh_study_child_confirm` under the ID CONS1). Copy it exactly.
- `[random]` when the questionnaire says to rotate or randomise the options.
- `[hint: text]` an interviewer instruction for that question that the questionnaire does not print in square brackets ("Read out the options", "Do not read"). No square brackets inside the hint; bracketed instructions stay in the text (see markers below).

Option lines contain only the option text, e.g. `Po`: no option numbers, letters, checkboxes or dotted leaders. Keep the questionnaire's order and keep options like "Don't know" / "Refused" when listed. Tags allowed on option lines:
- `[code: X]` when the questionnaire gives answer codes (Yes 1, No 0, Don't know -98, Refused -99, Other 96 …): copy each code exactly, on every option of that question. When the questionnaire states a general coding convention (e.g. "1 = Yes; 0 = No", "-98 = Don't know; -99 = Refused"), apply it to options that have no code of their own. Without any codes or convention, leave this tag out; options are then numbered 1, 2, 3, …
- `[exclusive]` on an option of a `[multiple]` question that cannot be combined with other answers ("None", "No one"). Codes -97, -98 and -99 are exclusive automatically.

An option asking the respondent to specify ("Other, specify", "Tjetër (specifiko)") must end its text with `____` (four underscores), before any tags; this creates the follow-up text field. Any underscore in an option triggers that, so never use underscores in other options.

Special answers for `[numeric]`, `[decimal]`, `[date]`, `[time]` and `[text]`: when the questionnaire lists answers such as "Don't know -98", "Refused -99" or "Nothing 0" next to the value, put them as option lines after the question, each with its `[code: X]`. If the questionnaire also lists the value itself as an option with a code (e.g. "0–7 days 1", "Amount in euros 1"), include that line too and add `[value]` to it. Only lines with `[code: X]` or `[value]` are read as special answers.

Placeholders in question text such as [STUDY CHILD NAME] (text the interviewer replaces) must use parentheses: (STUDY CHILD NAME).

Bracketed markers that the questionnaire prints in its text, such as [READ ALOUD], [LEXO ME ZË TË LARTË], [INTERVIEWER NOTE: do not read the options], [SHËNIM I INTERVISTUESIT], [SECTION] or [SEKSIONI], are shown to the interviewer in the form, so copy them with their brackets and wording exactly where they stand: in the question, note, option or section text, after the question number, in every language, e.g. `A5. [READ ALOUD] Now I will read some statements. || [LEXO ME ZË TË LARTË] Tani do t'ju lexoj disa pohime. [note]`. Do not move them into a `[hint: …]`, and do not drop them. A marker is text, not a tag: tags still go at the end of the line. When the marker is only a single word that is also a tag word (section, note, text, other, single, multiple, group, repeat, random, value, date, time …), write it with `!` after the opening bracket so it is not read as a tag: `[!SECTION]`, `[!NOTE]`; the form shows it as [SECTION], [NOTE]. A heading marked with such a marker becomes a `[section]` line that keeps the marker in its title, e.g. `[section] [!SECTION] Informed consent || [SEKSIONI] Pëlqimi i informuar`.

## Structure

- Modules or major parts: a line `[group] <module title>` where the module starts and `[end group]` where it ends.
- Sections inside a module (every heading that introduces a block of questions, e.g. "[SECTION] Informed consent", "Section B: Employment"): a line `[section] <section title>`, e.g. `[section] Informed consent || Pëlqimi i informuar`. Each section becomes a group in the form, labelled with its title, and ends automatically at the next `[section]`, `[group]`, `[end group]` or `[end repeat]` — never write an end tag for a section. The title is only the heading; read-aloud or introduction text under it goes on its own `[note]` line after the `[section]` line.
- A questionnaire without modules can use `[section]` lines alone.
- Questions the questionnaire asks once per household member, child or other item (a roster or loop): `[repeat: <number of the question that gives how many times>] <title>` before them and `[end repeat]` after them, e.g. `[repeat: ROST0] Household roster`. Leave out the reference when no question holds the count: `[repeat] Children`.

Any other untagged line that is not an option, matrix column, matrix row or note continuation is dropped with a warning, so do not output stray text such as instructions, titles or table headers.

## Routing

Filters are built in a later step from routing notes you copy next to the place they apply. Capture every routing instruction of the questionnaire this way; it is checked automatically:
- `[ask if: <condition>]` on a question, `[group]`, `[section]` or `[repeat]` line when the questionnaire says who is asked there ("Ask if: If CONS1 = Yes", "Filter: employed respondents", a [ROUTING] row over a block, "Only for treatment households"). Write the condition in the first language, with question IDs written as in your lines, e.g. `[ask if: If CONS1 = Yes]`. Copy it as written even when it is incomplete (e.g. "ROST3 >= ." with the number missing) and mention that in `notes`. Leave out notes that do not restrict who is asked ("Ask once", "Ask first", "Ask all study children", "Ask once per activity").
- `[skip: TARGET]` on an option line when choosing that option jumps to another question (">>END2", "Go to Q7"): TARGET is the question ID as written in your lines, or `END` for the end of the questionnaire. A jump printed after the option list that applies to several options goes on each of those options.
- No square brackets inside these tags; use parentheses.
Keep all questions the routing refers to.

Order: keep the questionnaire's order, with one exception: closing / end-of-interview questions (final interview status, reason not completed, recontact permission, closing thank-you) go last, at the end of the form, even when the document prints them earlier (e.g. right after the consent module because refusals jump there). Refusals and ineligible cases jump to them with `[skip: …]`, and a completed interview reaches them after the last module.

## Languages

If the questionnaire has its text in more than one language (parallel columns or repeated blocks), code all of them. Decide from the text actually present in the questions and options, not from statements about the document: a note saying a column is blank or pending translation may be outdated. Every language that has text for the questions is coded.
- The first line is `[languages: <language 1>, <language 2>, …]` with the English names of the languages, in the questionnaire's order, e.g. `[languages: English, Albanian]`.
- Every text — question, option, note, hint, group/repeat title, matrix column and row, scale labels — is written as `<language 1 text> || <language 2 text>`, in the declared order. Tags are written once, at the end of the line. The question number is written once, at the start.
- Leave out a language whose text is missing or only a placeholder (e.g. "[Serbian translation to be inserted]").
With a single language, there is no `[languages: …]` line and no `||`. Never translate text yourself; copy each language as written.

Remove answer blanks (____) and codes from question text.

## notes

In `notes`, list (in Albanian) anything the user should check: ambiguous question types, unclear numbering, preloaded lists or calculations that must be added manually, content you could not read. Leave it empty when there is nothing to flag.

## Example (two languages, codes and variable names)

[languages: English, Albanian]
[group] MODULE A: GENERAL INFORMATION || MODULI A: TË DHËNA TË PËRGJITHSHME
A1. How old are you? || Sa vjeç jeni? [numeric] [name: resp_age]
Don't know || Nuk e di [code: -98]
Refused || Refuzon [code: -99]
A2. What is your employment status? || Cili është statusi juaj i punësimit? [single] [name: emp_status] [hint: Read the options || Lexo opsionet] [ask if: If A1 >= 15]
Employed || I/e punësuar [code: 1]
Unemployed || I/e papunë [code: 2] [skip: A4]
Other, specify ____ || Tjetër, specifiko ____ [code: 96]
Refused || Refuzon [code: -99]
A3. Which media do you use for news? || Cilat media përdorni për lajme? [multiple] [random] [name: media_used]
Television || Televizion [code: 1]
Internet || Internet [code: 2]
None || Asnjë [code: 3] [exclusive]
Don't know || Nuk e di [code: -98]
A4. How much did you pay last month? || Sa keni paguar muajin e kaluar? [decimal] [name: paid_amount]
Amount in euros || Shuma në euro [code: 1] [value]
Nothing || Asgjë [code: 0]
Don't know || Nuk e di [code: -98]
A5. How satisfied are you with public services? || Sa të kënaqur jeni me shërbimet publike? [scale 1(Not at all || Aspak)-5(Very || Shumë)]
[end group]
[group] MODULE B: HOUSEHOLD || MODULI B: FAMILJA
[section] Household roster || Lista e familjes
[READ ALOUD] First, I would like to list everyone who lives here. || [LEXO ME ZË TË LARTË] Së pari, dua të listoj të gjithë ata që jetojnë këtu. [note]
B0. How many people live in this household? || Sa persona jetojnë në këtë familje? [numeric] [name: hh_size]
[repeat: B0] Household members || Anëtarët e familjes
B1. What is the name of the member? || Si quhet anëtari? [text] [name: member_name]
B2. Is (NAME) female or male? || A është (EMRI) femër apo mashkull? [single] [name: member_sex]
Female || Femër [code: 1]
Male || Mashkull [code: 2]
[end repeat]
[end group]"""

CONVERT_SCHEMA = {
    "type": "object",
    "properties": {
        "lines": {"type": "array", "items": {"type": "string"}},
        "notes": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["lines", "notes"],
    "additionalProperties": False,
}


def convert_questionnaire(client, source_blocks):
    content = source_blocks + [{"type": "text", "text": "Convert this questionnaire into the tagged line format."}]
    data, usage = _call_claude(client, CONVERT_SYSTEM, content, CONVERT_SCHEMA, max_tokens=128000,
                               effort=CONVERT_EFFORT)
    lines = [line.strip() for line in data["lines"] if line.strip()]
    if not lines:
        raise AIError("Claude nuk gjeti pyetje në dokument.")
    return lines, data["notes"], usage


# ---------------------------------------------------------------------------
# Step 1b: translation check of the tagged lines
# ---------------------------------------------------------------------------

LANG_SEP = "||"
TAG_RE = re.compile(r'\[([^\]]*)\]')
_TAG_WORDS = {"single", "multiple", "text", "string", "numeric", "decimal", "date", "time", "note", "other",
              "random", "exclusive", "value", "group", "end group", "end_group", "repeat", "end repeat",
              "end_repeat", "section", "seksion"}
_TAG_PATTERN = re.compile(r'(hint|name|code|skip|ask if|repeat|languages?)\s*:|matrix\s+(single|multiple)\s+\d'
                          r'|ranking\s+\d|scale\s*\d', flags=re.IGNORECASE)


def is_tag(tag):
    """True for the program's tags ([single], [hint: …], [scale 1-5] …). Any other bracketed text, such as
    [READ ALOUD] or [INTERVIEWER NOTE], is part of the question text and is kept in the form."""
    tag = tag.strip()
    return tag.lower() in _TAG_WORDS or _TAG_PATTERN.match(tag) is not None


def strip_tags(line):
    """The line without the program's tags; bracketed markers such as [READ ALOUD] stay."""
    text = TAG_RE.sub(lambda m: " " if is_tag(m.group(1)) else m.group(0), line)
    return re.sub(r'\s{2,}', ' ', text).strip()


def unescape_markers(text):
    """[!NOTE] is how a marker whose word is also a tag is written; in the form it reads [NOTE]."""
    return text.replace("[!", "[") if isinstance(text, str) else text


def declared_languages(lines):
    """Languages of a [languages: English, Albanian] line; [] for a single-language questionnaire."""
    for line in lines:
        m = re.match(r'^\s*\[languages?:\s*(.+?)\]\s*$', line, flags=re.IGNORECASE)
        if m:
            langs = [l.strip() for l in re.split(r'\|\||,', m.group(1)) if l.strip()]
            return langs if len(langs) > 1 else []
    return []


def _untranslated(text, languages):
    """Languages (after the first) whose part of a `a || b` text is missing, empty or a copy of the first."""
    parts = [p.strip() for p in text.split(LANG_SEP)]
    first = re.sub(r'^\w+[.)]\s+', '', parts[0]).strip().lower()   # the question number is only in the first part
    missing = []
    for k, lang in enumerate(languages[1:], start=1):
        part = parts[k] if k < len(parts) else ""
        # identical short texts ("Internet", "Facebook") are normal; identical sentences are not
        if not part or (part.lower() == first and len(first.split()) >= 3):
            missing.append(lang)
    return missing


def check_translations(lines):
    """Text lines of a multi-language questionnaire that are missing a language (or keep the first
    language's text in its place). Returns [{"index", "languages", "line"}]; [] with a single language."""
    languages = declared_languages(lines)
    if not languages:
        return []
    issues = []
    for i, line in enumerate(lines):
        missing = set()
        text = strip_tags(line)
        if text:
            missing.update(_untranslated(text, languages))
        for tag in TAG_RE.findall(line):
            m = re.match(r'\s*hint\s*:\s*(.+)', tag, flags=re.IGNORECASE)
            if m:
                missing.update(_untranslated(m.group(1), languages))
        if missing:
            issues.append({"index": i, "languages": [l for l in languages if l in missing], "line": line})
    return issues


def _structure(line):
    """What a translation fix must not change: the question number and the tags (without their texts)."""
    tags = [re.sub(r'\([^)]*\)', '', t).strip().lower() for t in TAG_RE.findall(line)
            if is_tag(t) and not re.match(r'\s*hint\s*:', t, flags=re.IGNORECASE)]
    tags = [t.split(":")[0] if LANG_SEP in t else t for t in tags]
    number = re.match(r'^\s*(\w+)[.)]\s', strip_tags(line))
    return tags, number.group(1) if number else None


TRANSLATION_SYSTEM = INPUT_NOTE + """

A multi-language questionnaire was converted into a line-based tagged format. In that format every text (question, option, note, hint, group/section title, matrix column and row) is written as `<language 1 text> || <language 2 text> || …` in the order of the `[languages: …]` line; tags are written once at the end of the line (at the start for `[group]`, `[section]` and `[repeat]` lines), and the question number only once, at the start.

An automatic check found lines where a language is missing or still holds the first language's text. For each listed line, return the complete corrected line with the text of every language filled in, taken from the questionnaire. Change nothing else: keep the question number, the first language's text, all tags and their order exactly as they are.

Copy each language's text as the questionnaire writes it. Only when the questionnaire truly has no text in that language for this item, translate it yourself, and list those lines in `notes` (in Albanian) so the user can have them checked.

A line whose text really is the same in every language (a brand name, a number) can be returned unchanged."""

TRANSLATION_SCHEMA = {
    "type": "object",
    "properties": {
        "fixes": {"type": "array", "items": {
            "type": "object",
            "properties": {"index": {"type": "integer"}, "line": {"type": "string"}},
            "required": ["index", "line"],
            "additionalProperties": False,
        }},
        "notes": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["fixes", "notes"],
    "additionalProperties": False,
}


def repair_translations(client, source_blocks, lines, rounds=2):
    """Fills in missing languages from the questionnaire and tests the lines again after every round.
    Returns (lines, remaining issues, notes, cost)."""
    lines = list(lines)
    languages = declared_languages(lines)
    issues = check_translations(lines)
    notes, cost = [], 0.0
    for _ in range(rounds):
        if not issues:
            break
        listed = "\n".join(f"{p['index']}: {p['line']}    <-- mungon: {', '.join(p['languages'])}" for p in issues)
        numbered = "\n".join(f"{i}: {line}" for i, line in enumerate(lines))
        content = _cached(source_blocks) + [{"type": "text", "text":
            f"Languages, in order: {', '.join(languages)}\n\n"
            f"<converted_lines>\n{numbered}\n</converted_lines>\n\n"
            f"Lines to fix (index: line <-- missing languages):\n<to_fix>\n{listed}\n</to_fix>"}]
        data, usage = _call_claude(client, TRANSLATION_SYSTEM, content, TRANSLATION_SCHEMA,
                                   max_tokens=64000, effort=CONVERT_EFFORT)
        cost += estimate_cost(usage)
        notes.extend(data["notes"])
        wanted = {p["index"] for p in issues}
        for fix in data["fixes"]:
            i, new = fix["index"], fix["line"].strip()
            # a fix that changes the number or the tags would change the form, not only its texts
            if i in wanted and new and _structure(new) == _structure(lines[i]):
                lines[i] = new
        issues = check_translations(lines)
    return lines, issues, notes, cost


def translation_warnings(issues, limit=30):
    """Messages for the lines that are still not translated."""
    msgs = [f"Rreshti {p['index'] + 1} nuk ka tekst në {', '.join(p['languages'])}: `{p['line'][:90]}`"
            for p in issues[:limit]]
    if len(issues) > limit:
        msgs.append(f"… dhe {len(issues) - limit} rreshta të tjerë pa përkthim.")
    return msgs


# ---------------------------------------------------------------------------
# Step 1c: second check of the notes (introductions, read-aloud text) and their place
# ---------------------------------------------------------------------------

NOTES_SYSTEM = INPUT_NOTE + """

A questionnaire was converted into a line-based tagged format. `[note]` lines hold text shown without an answer: introductions, consent text, read-aloud passages, section introductions; untagged lines right after a `[note]` line continue that note. `[group]` / `[section]` lines open a module / section, `[end group]` closes a module.

Check every note against the questionnaire, in two ways:
1. Place: a note must stand exactly where the questionnaire prints it. An introduction comes before the questions it introduces, right after the `[group]` / `[section]` line when it opens a module or section; a note placed after the questions it introduces, in another module, or at the end is misplaced. For each misplaced note give `move` with `index` (the index of its `[note]` line; its continuation lines move with it) and `before` (the index of the line it must stand right before, as numbered in the converted lines; the number of lines to put it at the very end).
2. Completeness: introductions, consent text or read-aloud passages of the questionnaire that are missing from the lines. For each give `insert` with `before` and `line`: the complete `[note]` line in the same format as the other lines (all languages separated by ` || ` when there is a `[languages: …]` line, the `[note]` tag at the end). Do not add interviewer instructions that belong to a single question, titles, or text already present.

Report only real problems; when every note is in place and complete, return empty lists. In `notes`, describe (in Albanian) each change in one short sentence."""

NOTES_SCHEMA = {
    "type": "object",
    "properties": {
        "move": {"type": "array", "items": {
            "type": "object",
            "properties": {"index": {"type": "integer"}, "before": {"type": "integer"}},
            "required": ["index", "before"],
            "additionalProperties": False,
        }},
        "insert": {"type": "array", "items": {
            "type": "object",
            "properties": {"before": {"type": "integer"}, "line": {"type": "string"}},
            "required": ["before", "line"],
            "additionalProperties": False,
        }},
        "notes": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["move", "insert", "notes"],
    "additionalProperties": False,
}


def _is_note(line):
    return any(t.strip().lower() == "note" for t in TAG_RE.findall(line))


def _note_block(lines, i):
    """Indexes of a note line and the untagged lines that continue it."""
    end = i + 1
    while end < len(lines) and not any(is_tag(t) for t in TAG_RE.findall(lines[end])):
        end += 1
    return list(range(i, end))


def _apply_note_fixes(lines, moves, inserts):
    """Moves note blocks and inserts missing notes; indexes refer to `lines` before any change.
    Returns (new lines, number of changes)."""
    items = list(enumerate(lines))              # (original index, line); inserted lines have index None
    changes = 0

    def position(orig):
        if orig >= len(lines):
            return len(items)
        return next((p for p, (o, _) in enumerate(items) if o == orig), None)

    for mv in moves:
        i, before = mv["index"], mv["before"]
        if not (0 <= i < len(lines)) or not _is_note(lines[i]) or not (0 <= before <= len(lines)):
            continue
        block = set(_note_block(lines, i))
        if before in block or before == max(block) + 1:
            continue                              # already there
        moved = [it for it in items if it[0] in block]
        items = [it for it in items if it[0] not in block]
        p = position(before)
        if p is None:
            items.extend(moved)                   # target vanished; keep the note rather than lose it
            continue
        items[p:p] = moved
        changes += 1
    for ins in inserts:
        line, before = ins["line"].strip(), ins["before"]
        if not line or not _is_note(line) or not (0 <= before <= len(lines)):
            continue
        p = position(before)
        items.insert(len(items) if p is None else p, (None, line))
        changes += 1
    return [line for _, line in items], changes


def review_notes(client, source_blocks, lines, rounds=2):
    """Second check of the notes against the questionnaire: misplaced notes are moved, missing ones added,
    and the result is checked again until nothing changes. Returns (lines, notes, cost)."""
    lines = list(lines)
    notes, cost = [], 0.0
    for _ in range(rounds):
        numbered = "\n".join(f"{i}: {line}" for i, line in enumerate(lines))
        content = _cached(source_blocks) + [{"type": "text", "text":
            f"<converted_lines>\n{numbered}\n</converted_lines>\n\nCheck the notes of the converted lines."}]
        data, usage = _call_claude(client, NOTES_SYSTEM, content, NOTES_SCHEMA,
                                   max_tokens=32000, effort=CONVERT_EFFORT)
        cost += estimate_cost(usage)
        lines, changes = _apply_note_fixes(lines, data["move"], data["insert"])
        if not changes:
            break
        notes.extend(data["notes"])
    return lines, notes, cost


# ---------------------------------------------------------------------------
# XLSForm read / write / describe
# ---------------------------------------------------------------------------

def read_form(data):
    """Reads an XLSForm into {sheet_name: DataFrame of strings}."""
    try:
        sheets = pd.read_excel(BytesIO(data), sheet_name=None, dtype=str, keep_default_na=False)
    except Exception as e:
        raise AIError(f"Skedari XLS nuk mund të lexohet: {e}")
    if _sheet_key(sheets, "survey") is None:
        raise AIError("Skedari nuk ka fletën 'survey' – nuk duket si formular XLSForm.")
    survey_key = _sheet_key(sheets, "survey")
    survey = sheets[survey_key]
    survey.columns = [str(c).strip() for c in survey.columns]
    if "relevant" not in survey.columns:
        survey["relevant"] = ""
    return sheets


def write_form(sheets):
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.mask(df.eq("")).to_excel(writer, sheet_name=name, index=False)
    return buf.getvalue()


def _sheet_key(sheets, wanted):
    return next((k for k in sheets if str(k).strip().lower() == wanted), None)


def _label_column(df):
    cols = [c for c in df.columns if str(c).lower() == "label"]
    cols += [c for c in df.columns if str(c).lower().startswith("label")]
    return cols[0] if cols else None


def _choices_by_list(sheets):
    key = _sheet_key(sheets, "choices")
    if key is None:
        return {}
    df = sheets[key]
    label_col = _label_column(df)
    lists = {}
    for _, r in df.iterrows():
        list_name = str(r.get("list_name", "")).strip()
        name = str(r.get("name", "")).strip()
        if list_name and name:
            lists.setdefault(list_name, []).append((name, str(r.get(label_col, "")) if label_col else ""))
    return lists


def describe_form(sheets, max_choices=60, routing=None):
    """Compact text view of the survey rows and the choice lists they use, with the routing captured from the
    questionnaire ([ask if] / [skip] tags) next to the rows it belongs to."""
    survey = sheets[_sheet_key(sheets, "survey")]
    label_col = _label_column(survey)
    ask_if = (routing or {}).get("ask_if", {})
    skips = {}
    for sk in (routing or {}).get("skips", []):
        skips.setdefault(sk["source"], []).append(f"{sk['code']} >> {sk['target']}")
    rows, used_lists = [], []
    for _, r in survey.iterrows():
        q_type = str(r.get("type", "")).strip()
        if not q_type:
            continue
        parts = [q_type, str(r.get("name", "")).strip()]
        label = str(r.get(label_col, "")).strip().replace("\n", " ") if label_col else ""
        if label:
            parts.append(label[:200])
        rel = str(r.get("relevant", "")).strip()
        if rel:
            parts.append(f"relevant: {rel}")
        name = str(r.get("name", "")).strip()
        if name in ask_if:
            parts.append(f"ask if: {ask_if[name]}")
        if name in skips:
            parts.append("skips: " + ", ".join(skips[name]))
        rows.append(" | ".join(parts))
        tokens = q_type.split()
        if tokens[0].lower() in ("select_one", "select_multiple") and len(tokens) > 1:
            used_lists.append(tokens[1])

    lists = _choices_by_list(sheets)
    choice_lines = []
    for list_name in dict.fromkeys(used_lists):
        items = lists.get(list_name, [])
        shown = "; ".join(f"{n}={lbl}" for n, lbl in items[:max_choices])
        if len(items) > max_choices:
            shown += f"; … ({len(items) - max_choices} more)"
        choice_lines.append(f"{list_name}: {shown}")

    return ("<survey>\n(type | name | label | current relevant | routing from the questionnaire)\n" + "\n".join(rows) + "\n</survey>\n\n"
            "<choices>\n(list: name=label; …)\n" + "\n".join(choice_lines) + "\n</choices>")


# ---------------------------------------------------------------------------
# Step 2: skip logic
# ---------------------------------------------------------------------------

FILTER_SYSTEM = INPUT_NOTE + """

You add skip logic to a KoboToolbox XLSForm that was built from a questionnaire. You get the original questionnaire plus the form's survey rows and choice lists. Find every routing or filter instruction in the questionnaire and express it as XLSForm `relevant` expressions on the right rows.

How to express filters:
- Put the condition on the row that is shown conditionally. "Ask Q5 only if Q4 = Yes" → target Q5, relevant `${Q4} = '1'` (using Q4's choice name for "Yes").
- "Go to" / "skip to" instructions: each question between the source question and the destination gets the condition under which it is asked. Prefer listing the answers that continue (`${Q3} = '1' or ${Q3} = '3'`) over `!=`, because `${Q3} != '2'` is also true when Q3 itself was hidden.
- A section-level filter ("Section C only for employed respondents") goes on every question of that section.
- When several instructions apply to one row, combine them into one expression with `and` / `or` and parentheses.
- Use the exact row names and choice names from the form (choice names, not labels). select_one: `${Q} = 'name'`. select_multiple: `selected(${Q}, 'name')`, negated with `not(selected(${Q}, 'name'))`. integer: `${Q} >= 18` with an unquoted number.
- Matrix and ranking questions are groups (`…_group` rows); target the group row, not its rows.
- Modules are groups (`group_N` rows), sections are groups too (`section_N` rows, labelled with the section title) and loops are repeats (`repeat_N` rows). A condition for a whole module, section or loop goes on that row instead of on each question inside it.
- A question with a value and special answers has two rows: `X_type` (select_one: value / Don't know / Refused …) and `X` (the value, shown when the value option is chosen). Conditions on the number use `${X}` (e.g. `${X} >= 18`); conditions on "Don't know" use `${X_type}`. The target for "ask X" is `X_type`.
- Answer codes can be 0 or negative (No = '0', Don't know = '-98'); always use the choice names listed in the form.
- Rows that already have a relevant (automatic "specify" fields) need no changes unless the questionnaire adds a condition; if you target one, give the complete expression, as `relevant` replaces the current value.
- Rows may carry routing copied from the questionnaire: `ask if: …` (who is asked) and `skips: CODE >> TARGET` (choosing that answer jumps to TARGET, a question ID; END is the end of the questionnaire). Implement every one of them. Automated tests check that each `ask if` row, or a group around it, has a filter, and that for each skip every question between the source and TARGET is hidden when that answer is chosen. A routing note you cannot implement (incomplete, like "ROST3 >= .", or about a question that is not in the form) is left out and explained in `notes`.
- Only reference rows that come before the target.
- Add only filters the questionnaire states or clearly implies (e.g. a "Why not?" follow-up to a yes/no question). Do not invent logic.
- `instruction`: a short quote (at most about 15 words) of the questionnaire text the filter is based on.
- `notes` (in Albanian): instructions you could not map to the form (e.g. they refer to a question that is not in the form) and anything the user should verify. Leave empty when there is nothing to flag."""

FILTER_SCHEMA = {
    "type": "object",
    "properties": {
        "filters": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "target": {"type": "string"},
                    "relevant": {"type": "string"},
                    "instruction": {"type": "string"},
                },
                "required": ["target", "relevant", "instruction"],
                "additionalProperties": False,
            },
        },
        "notes": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["filters", "notes"],
    "additionalProperties": False,
}


def generate_filters(client, source_blocks, sheets, failed_tests=None, routing=None):
    """Asks Claude for relevant expressions. With failed_tests, asks only for fixes of those rows."""
    text = describe_form(sheets, routing=routing)
    if failed_tests:
        text += ("\n\nThe filters below failed automated tests. Return corrected filters only for these rows "
                 "(the complete new relevant for each). If a filter should be removed, return an empty relevant. "
                 "A routing problem can also be fixed on a group around the row or on an earlier row it depends on.\n"
                 + format_issues(failed_tests))
    else:
        text += "\n\nAdd the skip logic for this form."
    content = _cached(source_blocks) + [{"type": "text", "text": text}]
    data, usage = _call_claude(client, FILTER_SYSTEM, content, FILTER_SCHEMA, max_tokens=64000)
    return data["filters"], data["notes"], usage


def apply_filters(sheets, filters, combine_existing=True):
    """Writes relevant expressions into the survey sheet.

    With combine_existing, a row that already has a relevant keeps it and the new condition is
    added with `and` (used for the first pass, where only automatic "specify" fields have one).
    Returns (applied, rejected) lists of dicts for display.
    """
    survey = sheets[_sheet_key(sheets, "survey")]
    names = survey["name"].astype(str).str.strip() if "name" in survey.columns else pd.Series(dtype=str)
    applied, rejected = [], []
    for f in filters:
        target = f["target"].strip()
        new_rel = f["relevant"].strip()
        idx = names.index[names == target]
        if len(idx) == 0:
            rejected.append({**f, "reason": "Rreshti nuk ekziston në formular"})
            continue
        i = idx[0]
        if str(survey.at[i, "type"]).strip().lower().replace(" ", "_") in ("end_group", "end_repeat"):
            rejected.append({**f, "reason": "Filtri nuk mund të vendoset në fund të grupit"})
            continue
        old_rel = str(survey.at[i, "relevant"]).strip()
        if combine_existing and old_rel and new_rel and new_rel != old_rel:
            new_rel = f"({old_rel}) and ({new_rel})"
        survey.at[i, "relevant"] = new_rel
        applied.append({**f, "relevant": new_rel})
    return applied, rejected


# ---------------------------------------------------------------------------
# Filter check (existing XLSForm)
# ---------------------------------------------------------------------------

REVIEW_SYSTEM = INPUT_NOTE + """

You review the skip logic (the `relevant` column) of a KoboToolbox XLSForm. You get the original questionnaire, the form's survey rows and choice lists, and the results of automated tests on every relevant expression. The questionnaire is the reference: the form's filters must implement its routing.

Report problems that change who sees which question:
- wrong choice names or values, wrong operators (select_multiple needs `selected()`), syntax errors
- conditions that contradict the questionnaire's routing, or routing in the questionnaire with no filter in the form
- follow-up questions shown to everyone although their label depends on an earlier answer (e.g. "Why not?")
- references to later questions, conditions that can never be true or never hide anything
Confirm or explain each automated test finding. Do not report style preferences or equivalent rewrites.

For each issue: `name` is the row to change; `action` is "replace" (set `suggested_relevant` as the complete new expression), "remove" (clear the relevant) or "none" (only a remark); `problem` explains it in one or two sentences in Albanian. Use the exact row and choice names from the form. `summary` is a short overall assessment in Albanian."""

REVIEW_SCHEMA = {
    "type": "object",
    "properties": {
        "summary": {"type": "string"},
        "issues": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "name": {"type": "string"},
                    "severity": {"type": "string", "enum": ["error", "warning", "info"]},
                    "problem": {"type": "string"},
                    "action": {"type": "string", "enum": ["replace", "remove", "none"]},
                    "suggested_relevant": {"type": "string"},
                },
                "required": ["name", "severity", "problem", "action", "suggested_relevant"],
                "additionalProperties": False,
            },
        },
    },
    "required": ["summary", "issues"],
    "additionalProperties": False,
}


def review_filters(client, sheets, test_report, source_blocks):
    text = describe_form(sheets) + "\n\n<automated_tests>\n" + (format_issues(test_report["issues"]) or "No findings.") \
        + "\n</automated_tests>\n\nReview the skip logic of this form against the routing in the questionnaire."
    content = source_blocks + [{"type": "text", "text": text}]
    data, usage = _call_claude(client, REVIEW_SYSTEM, content, REVIEW_SCHEMA, max_tokens=64000)
    return data, usage


def apply_suggestions(sheets, issues):
    filters = [{"target": i["name"], "relevant": i["suggested_relevant"] if i["action"] == "replace" else "",
                "instruction": i["problem"]}
               for i in issues if i["action"] in ("replace", "remove")]
    return apply_filters(sheets, filters, combine_existing=False)


def format_issues(issues):
    return "\n".join(f"- {i['name']} [{i['severity']}] relevant `{i.get('relevant', '')}`: {i['message']}"
                     for i in issues)


# ---------------------------------------------------------------------------
# Offline tests for relevant expressions
# ---------------------------------------------------------------------------

class RelevantSyntaxError(Exception):
    pass


class Unsupported(Exception):
    pass


_TOKEN_RE = re.compile(r"""\s*(?:
    (?P<var>\$\{\s*([A-Za-z_][\w.\-]*)\s*\})
  | (?P<str>'[^']*'|"[^"]*")
  | (?P<num>\d+(?:\.\d+)?|\.\d+)
  | (?P<op>!=|<=|>=|=|<|>|\+|-|\*|\(|\)|,)
  | (?P<name>[A-Za-z_][\w\-]*(?::[A-Za-z_][\w\-]*)?)
  | (?P<dot>\.\.|\.)
  | (?P<path>/)
)""", re.X)


def _tokenize(expr):
    tokens, pos = [], 0
    expr = expr.rstrip()
    while pos < len(expr):
        m = _TOKEN_RE.match(expr, pos)
        if not m or m.end() == pos:
            raise RelevantSyntaxError(f"karakter i panjohur '{expr[pos:].strip()[:10]}'")
        kind = m.lastgroup
        if kind == "var":
            tokens.append(("var", m.group(2)))
        elif kind == "str":
            tokens.append(("str", m.group(kind)[1:-1]))
        else:
            tokens.append((kind, m.group(kind)))
        pos = m.end()
    return tokens


class _Parser:
    COMPARE = ("=", "!=", "<", "<=", ">", ">=")

    def __init__(self, tokens):
        self.tokens, self.i = tokens, 0

    def peek(self):
        return self.tokens[self.i] if self.i < len(self.tokens) else (None, None)

    def take(self):
        tok = self.peek()
        self.i += 1
        return tok

    def expect(self, value):
        tok = self.take()
        if tok[1] != value:
            raise RelevantSyntaxError(f"pritej '{value}' por u gjet '{tok[1] or 'fundi i shprehjes'}'")

    def parse(self):
        if not self.tokens:
            raise RelevantSyntaxError("shprehja është bosh")
        node = self.or_()
        if self.i != len(self.tokens):
            raise RelevantSyntaxError(f"'{self.peek()[1]}' e papritur")
        return node

    def or_(self):
        node = self.and_()
        while self.peek() == ("name", "or"):
            self.take()
            node = ("or", node, self.and_())
        return node

    def and_(self):
        node = self.cmp()
        while self.peek() == ("name", "and"):
            self.take()
            node = ("and", node, self.cmp())
        return node

    def cmp(self):
        node = self.add()
        while self.peek()[0] == "op" and self.peek()[1] in self.COMPARE:
            op = self.take()[1]
            node = ("cmp", op, node, self.add())
        return node

    def add(self):
        node = self.mul()
        while self.peek() in (("op", "+"), ("op", "-")):
            op = self.take()[1]
            node = ("arith", op, node, self.mul())
        return node

    def mul(self):
        node = self.unary()
        while self.peek() in (("op", "*"), ("name", "div"), ("name", "mod")) or self.peek()[0] == "path":
            op = self.take()[1]
            rhs = self.unary()
            # XPath paths (../name, position(..)) are valid in ODK but not evaluated by these tests
            node = ("dot",) if op == "/" else ("arith", op, node, rhs)
        return node

    def unary(self):
        if self.peek() == ("op", "-"):
            self.take()
            return ("neg", self.unary())
        return self.primary()

    def primary(self):
        kind, val = self.take()
        if kind == "num":
            return ("num", float(val))
        if kind == "str":
            return ("str", val)
        if kind == "var":
            return ("var", val)
        if kind == "dot":
            return ("dot",)
        if (kind, val) == ("op", "("):
            node = self.or_()
            self.expect(")")
            return node
        if kind == "name" and self.peek() == ("op", "("):
            self.take()
            args = []
            if self.peek() != ("op", ")"):
                args.append(self.or_())
                while self.peek() == ("op", ","):
                    self.take()
                    args.append(self.or_())
            self.expect(")")
            return ("call", val, args)
        if kind is None:
            raise RelevantSyntaxError("shprehja përfundon papritur")
        raise RelevantSyntaxError(f"'{val}' e papritur")


def parse_relevant(expr):
    return _Parser(_tokenize(expr)).parse()


def _walk(node):
    yield node
    for child in node[1:]:
        if isinstance(child, tuple):
            yield from _walk(child)
        elif isinstance(child, list):
            for c in child:
                yield from _walk(c)


def _vars(node):
    return {n[1] for n in _walk(node) if n[0] == "var"}


def _to_str(v):
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float):
        if math.isnan(v):
            return "NaN"
        return str(int(v)) if v.is_integer() else str(v)
    return v


def _to_num(v):
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if isinstance(v, float):
        return v
    try:
        return float(v.strip())
    except ValueError:
        return float("nan")


def _to_bool(v):
    if isinstance(v, bool):
        return v
    if isinstance(v, float):
        return v != 0 and not math.isnan(v)
    return len(v) > 0


def _literal_pairs(node):
    """(var, literal) pairs from `${v} = 'x'`, `${v} != 'x'` and `selected(${v}, 'x')`."""
    pairs = []
    for n in _walk(node):
        if n[0] == "cmp" and n[1] in ("=", "!="):
            a, b = n[2], n[3]
            if b[0] == "var":
                a, b = b, a
            if a[0] == "var" and b[0] in ("str", "num"):
                pairs.append((a[1], _to_str(b[1]), n[1]))
        elif n[0] == "call" and n[1] == "selected" and len(n[2]) == 2:
            a, b = n[2]
            if a[0] == "var" and b[0] in ("str", "num"):
                pairs.append((a[1], _to_str(b[1]), "selected"))
    return pairs


def _compare(op, a, b):
    if op in ("=", "!="):
        if isinstance(a, bool) or isinstance(b, bool):
            eq = _to_bool(a) == _to_bool(b)
        elif isinstance(a, float) or isinstance(b, float):
            eq = _to_num(a) == _to_num(b)
        else:
            eq = a == b
        return eq if op == "=" else not eq
    x, y = _to_num(a), _to_num(b)
    return {"<": x < y, "<=": x <= y, ">": x > y, ">=": x >= y}[op]


def _apply_fn(name, vals):
    if name == "selected":
        return _to_str(vals[1]) in _to_str(vals[0]).split()
    if name == "not":
        return not _to_bool(vals[0])
    if name == "count-selected":
        return float(len(_to_str(vals[0]).split()))
    if name == "string-length":
        return float(len(_to_str(vals[0])))
    if name == "true":
        return True
    if name == "false":
        return False
    if name == "boolean":
        return _to_bool(vals[0])
    if name == "number":
        return _to_num(vals[0])
    if name == "string":
        return _to_str(vals[0])
    if name == "int":
        n = _to_num(vals[0])
        return float(math.trunc(n)) if not math.isnan(n) else n
    if name == "if":
        return vals[1] if _to_bool(vals[0]) else vals[2]
    if name == "coalesce":
        return vals[0] if _to_str(vals[0]) != "" else vals[1]
    if name == "contains":
        return _to_str(vals[1]) in _to_str(vals[0])
    if name == "starts-with":
        return _to_str(vals[0]).startswith(_to_str(vals[1]))
    if name == "ends-with":
        return _to_str(vals[0]).endswith(_to_str(vals[1]))
    if name == "concat":
        return "".join(_to_str(v) for v in vals)
    raise Unsupported(name)


def _eval(node, env):
    kind = node[0]
    if kind in ("num", "str"):
        return node[1]
    if kind == "var":
        return env.get(node[1], "")
    if kind == "dot":
        raise Unsupported(".")
    if kind == "neg":
        return -_to_num(_eval(node[1], env))
    if kind == "or":
        return _to_bool(_eval(node[1], env)) or _to_bool(_eval(node[2], env))
    if kind == "and":
        return _to_bool(_eval(node[1], env)) and _to_bool(_eval(node[2], env))
    if kind == "cmp":
        return _compare(node[1], _eval(node[2], env), _eval(node[3], env))
    if kind == "arith":
        return _arith(node[1], _eval(node[2], env), _eval(node[3], env))
    if kind == "call":
        return _apply_fn(node[1], [_eval(a, env) for a in node[2]])
    raise Unsupported(kind)


def _arith(op, x, y):
    x, y = _to_num(x), _to_num(y)
    if op == "+":
        return x + y
    if op == "-":
        return x - y
    if op == "*":
        return x * y
    if y == 0 or math.isnan(y):
        return float("nan")
    return x / y if op == "div" else math.fmod(x, y)


def _truthy(node, env):
    return _to_bool(_eval(node, env))


_UNKNOWN = object()   # three-valued logic: an answer that is not known


def _truth3(v):
    return v if v is _UNKNOWN else _to_bool(v)


def _eval3(node, env):
    """Like _eval, but a variable missing from env is unknown; the result is _UNKNOWN when it depends on one."""
    kind = node[0]
    if kind in ("num", "str"):
        return node[1]
    if kind == "var":
        return env.get(node[1], _UNKNOWN)
    if kind in ("and", "or"):
        a, b = _truth3(_eval3(node[1], env)), _truth3(_eval3(node[2], env))
        decisive = kind == "or"            # True decides an "or", False decides an "and"
        if a is decisive or b is decisive:
            return decisive
        return _UNKNOWN if (a is _UNKNOWN or b is _UNKNOWN) else not decisive
    if kind == "call" and node[1] == "not" and len(node[2]) == 1:
        v = _truth3(_eval3(node[2][0], env))
        return v if v is _UNKNOWN else not v
    if kind in ("call", "cmp", "arith", "neg"):
        args = node[2] if kind == "call" else node[2:4] if kind != "neg" else node[1:2]
        vals = [_eval3(a, env) for a in args]
        if any(v is _UNKNOWN for v in vals):
            return _UNKNOWN
        try:
            if kind == "call":
                return _apply_fn(node[1], vals)
            if kind == "cmp":
                return _compare(node[1], *vals)
            if kind == "arith":
                return _arith(node[1], *vals)
            return -_to_num(vals[0])
        except (Unsupported, IndexError):
            return _UNKNOWN
    return _UNKNOWN


def _implied_answers(conds):
    """Answers that must hold when all `conds` are true: the `${v} = 'x'` parts of their `and` chains."""
    out, stack = {}, list(conds)
    while stack:
        n = stack.pop()
        if n[0] == "and":
            stack += [n[1], n[2]]
        elif n[0] == "cmp" and n[1] == "=":
            a, b = n[2], n[3]
            if b[0] == "var":
                a, b = b, a
            if a[0] == "var" and b[0] in ("str", "num"):
                out[a[1]] = _to_str(b[1])
    return out


def _hidden_after(rows, conditions, source, code, by_name=None):
    """Names of the rows that are certainly hidden when `source` has answer `code` (hidden rows count as empty,
    so chains of filters are followed). Answers implied by the source being asked at all (e.g. it is only asked
    when EMP1..EMP4 = No) are known too; every other answer is unknown."""
    env = {source["name"]: code}
    pending = [source]
    while pending:   # the source is shown, so its filters hold - and those of the answers they imply
        row = pending.pop()
        for var, value in _implied_answers(conditions(row)).items():
            if var not in env:
                env[var] = value
                if by_name and var in by_name:
                    pending.append(by_name[var])
    hidden = set()
    for r in rows:
        if r["name"] in env:
            continue
        shown = True
        for c in conditions(r):
            t = _truth3(_eval3(c, env))
            if t is False:
                shown = False
                break
        if not shown:
            hidden.add(r["name"])
            if r["name"]:
                env[r["name"]] = ""
    return hidden


MAX_COMBINATIONS = 50000

SEVERITY_LABELS = {"error": "Gabim", "warning": "Paralajmërim", "info": "Info"}


AUTO_ROWS = ("start", "end", "GPS", "Anketuesi_ja", "emri_mbiemri", "numri_telefonit")
NON_QUESTION_TYPES = ("begin_group", "begin_repeat", "begin", "note", "calculate", "start", "end")


def test_form(sheets, routing=None):
    """Tests every relevant expression in the survey. With routing (captured from the questionnaire), also
    checks that each [ask if] has a filter and that each [skip] hides the questions it jumps over.
    Returns {"issues", "tested", "errors", "warnings"}; routing issues have kind "routing"."""
    survey = sheets[_sheet_key(sheets, "survey")]
    lists = _choices_by_list(sheets)

    rows, by_name, group_stack = [], {}, []
    for i, r in survey.iterrows():
        q_type = str(r.get("type", "")).strip()
        if not q_type:
            continue
        norm = q_type.lower().replace(" ", "_")
        if norm in ("end_group", "end_repeat"):
            if group_stack:
                group_stack.pop()
            continue
        tokens = q_type.split()
        row = {
            "excel_row": i + 2,
            "pos": len(rows),
            "name": str(r.get("name", "")).strip(),
            "base": tokens[0].lower(),
            "list": tokens[1] if len(tokens) > 1 else None,
            "relevant": str(r.get("relevant", "")).strip(),
            "required": str(r.get("required", "")).strip().lower() in ("true", "yes", "1", "true()"),
            "groups": list(group_stack),
            "ast": None,
        }
        rows.append(row)
        if row["name"]:
            by_name.setdefault(row["name"], row)
        if norm in ("begin_group", "begin_repeat"):
            group_stack.append(row["name"])

    issues = []

    def add(row, severity, message, kind=None):
        issues.append({"row": row["excel_row"], "name": row["name"], "severity": severity,
                       "relevant": row["relevant"], "message": message, "kind": kind})

    filtered = [r for r in rows if r["relevant"]]
    for row in filtered:
        try:
            row["ast"] = parse_relevant(row["relevant"])
        except RelevantSyntaxError as e:
            add(row, "error", f"Gabim sintakse në filtër: {e}.")

    # Static checks
    testable = []
    for row in filtered:
        if row["ast"] is None:
            continue
        ok = True
        for v in sorted(_vars(row["ast"])):
            ref = by_name.get(v)
            if ref is None:
                add(row, "error", f"Pyetja '{v}' e përmendur në filtër nuk ekziston në formular.")
                ok = False
            elif ref is row:
                add(row, "error", "Filtri i referohet vetë kësaj pyetjeje.")
                ok = False
            elif ref["pos"] > row["pos"]:
                add(row, "warning", f"Filtri i referohet pyetjes '{v}' që vjen më vonë në formular.")
        for v, literal, op in _literal_pairs(row["ast"]):
            ref = by_name.get(v)
            if ref is None or ref["base"] not in ("select_one", "select_multiple"):
                continue
            names = [n for n, _ in lists.get(ref["list"], [])]
            # '' is "not answered", a valid comparison rather than an answer code
            if names and literal != "" and literal not in names:
                add(row, "error", f"Vlera '{literal}' nuk ekziston te opsionet e pyetjes '{v}' "
                                  f"(opsionet: {', '.join(names[:15])}{'…' if len(names) > 15 else ''}).")
                ok = False
            if ref["base"] == "select_multiple" and op in ("=", "!="):
                add(row, "warning", f"'{v}' lejon disa përgjigje: përdor selected(${{{v}}}, '{literal}') "
                                    f"në vend të '{op}'.")
        if ok:
            testable.append(row)

    # Simulation over answer combinations
    def conditions(row):
        """Parsed conditions that decide whether the row is shown (enclosing groups + own)."""
        conds = [by_name[g]["ast"] for g in row["groups"] if g in by_name and by_name[g]["ast"] is not None]
        if row["ast"] is not None:
            conds.append(row["ast"])
        return conds

    for row in testable:
        result = _simulate(row, conditions, by_name, lists)
        if result == "never":
            add(row, "error", "Pyetja nuk shfaqet asnjëherë: filtri nuk plotësohet me asnjë kombinim përgjigjesh.")
        elif result == "always":
            add(row, "warning", "Filtri plotësohet gjithmonë – nuk e fsheh asnjëherë pyetjen.")
        elif isinstance(result, tuple) and result[0] == "unsupported":
            add(row, "info", f"Filtri nuk u testua automatikisht (funksioni '{result[1]}' nuk mbështetet nga testi).")
        elif result == "too_many":
            add(row, "info", "Filtri ka shumë kombinime përgjigjesh për t'u testuar automatikisht.")

    if routing:
        _check_routing(routing, rows, by_name, conditions, add)

    issues.sort(key=lambda x: x["row"])
    return {
        "issues": issues,
        "tested": len(filtered),
        "errors": sum(1 for i in issues if i["severity"] == "error"),
        "warnings": sum(1 for i in issues if i["severity"] == "warning"),
    }


def _check_routing(routing, rows, by_name, conditions, add):
    """Coverage of the questionnaire's routing: every [ask if] needs a filter (on the row or a group around it)
    and every [skip] must hide all questions between the source and its target for that answer."""
    for name, text in routing.get("ask_if", {}).items():
        row = by_name.get(name)
        if row is not None and not row["relevant"] and not conditions(row):
            add(row, "warning", f"Pyetësori thotë '{text}', por kjo pyetje nuk ka filtër.", kind="routing")

    refs = routing.get("refs", {})
    simple = lambda ref: re.sub(r'[^a-z0-9]', '', ref.lower())
    refs_simple = {simple(k): v for k, v in refs.items()}
    for skip in routing.get("skips", []):
        source = by_name.get(skip["source"])
        if source is None:
            continue
        target = skip["target"].strip()
        if simple(target) in ("end", "fund", "endofinterview", "endofquestionnaire") and simple(target) not in refs_simple:
            end_pos = len(rows)
        else:
            target_row = refs.get(target) or refs_simple.get(simple(target))
            if target_row not in by_name:
                add(source, "info", f"Kalimi '>> {target}' pas përgjigjes '{skip['code']}' nuk u gjet në formular; "
                                    f"kontrolloje manualisht.", kind="routing")
                continue
            end_pos = by_name[target_row]["pos"]
        if end_pos <= source["pos"]:
            continue   # a jump back (loop) is not a filter
        hidden = _hidden_after(rows, conditions, source, skip["code"], by_name)
        # the source's own "specify" field (Other, specify ____ >> X) is asked before the jump
        own_specify = lambda r: (r["base"] == "text" and r["name"].startswith(source["name"] + "_")
                                 and f"${{{source['name']}}}" in r["relevant"])
        missed = [r for r in rows[source["pos"] + 1:end_pos]
                  if r["name"] and r["name"] not in hidden and r["name"] not in AUTO_ROWS
                  and r["base"] not in NON_QUESTION_TYPES and not own_specify(r)]
        for r in missed[:15]:
            add(r, "error", f"Pyetësori thotë: nëse {source['name']} = '{skip['code']}', kalo te {target}. "
                            f"Kjo pyetje mund të shfaqet ende pas asaj përgjigjeje.", kind="routing")
        if len(missed) > 15:
            add(source, "error", f"Edhe {len(missed) - 15} pyetje të tjera nuk anashkalohen kur {source['name']} = "
                                 f"'{skip['code']}' (kalo te {target}).", kind="routing")


def _domain(ref, lists, numbers, strings, small=False):
    base = ref["base"]
    if base == "select_one":
        values = [n for n, _ in lists.get(ref["list"], [])]
    elif base == "select_multiple":
        names = [n for n, _ in lists.get(ref["list"], [])]
        values = list(names)
        if not small and len(names) <= 10:
            values += [" ".join(p) for p in itertools.combinations(names, 2)]
        elif len(names) > 2:
            # long lists: a few multi-answer combinations so count-selected() > 1 can be true
            values += [" ".join(names[:2]), " ".join(names[:3])]
    elif base in ("integer", "decimal", "range"):
        # literals ± 1 plus small values, so comparisons between two numeric questions can go both ways
        values = [_to_str(float(n + d)) for n in numbers for d in (-1, 0, 1)] + ["0", "1", "2"]
    else:
        values = strings + ["tekst"]
    if not ref["required"] or not values:
        values.append("")
    return list(dict.fromkeys(values))


def _simulate(target, conditions, by_name, lists):
    """Returns "ok", "never", "always", "too_many" or ("unsupported", fn)."""
    target_conds = conditions(target)
    direct = set().union(*(_vars(c) for c in target_conds)) - {target["name"]}

    closure, stack = set(), list(direct)
    while stack:
        v = stack.pop()
        if v in closure or v not in by_name or v == target["name"]:
            continue
        closure.add(v)
        for c in conditions(by_name[v]):
            stack.extend(_vars(c))

    all_nodes = [n for c in target_conds for n in _walk(c)]
    for v in closure:
        for c in conditions(by_name[v]):
            all_nodes.extend(_walk(c))
    numbers = sorted({n[1] for n in all_nodes if n[0] == "num"})
    strings = sorted({n[1] for n in all_nodes if n[0] == "str"})

    # Full simulation (hidden questions become empty) when small enough, otherwise direct references only
    for variables, hide, small in ((closure, True, False), (direct, False, False), (direct, False, True)):
        order = sorted((v for v in variables if v in by_name), key=lambda v: by_name[v]["pos"])
        domains = [_domain(by_name[v], lists, numbers, strings, small) for v in order]
        if math.prod(len(d) for d in domains) <= MAX_COMBINATIONS:
            break
    else:
        return "too_many"

    shown, own_false = False, False
    own = target["ast"]
    try:
        for combo in itertools.product(*domains):
            env = {}
            for v, value in zip(order, combo):
                if hide and not all(_truthy(c, env) for c in conditions(by_name[v])):
                    env[v] = ""
                else:
                    env[v] = value
            own_ok = _truthy(own, env)
            if not own_ok:
                own_false = True
            elif all(_truthy(c, env) for c in target_conds):
                shown = True
            if shown and own_false:
                return "ok"
    except Unsupported as e:
        return ("unsupported", str(e))

    if not shown:
        return "never"
    return "always"
